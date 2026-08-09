from __future__ import annotations

import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from excel_privacy_cleaner.excel_processor import ExcelPrivacyProcessor, ProcessingOptions  # noqa: E402


def assert_true(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def create_simple_xlsx(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "連絡先"
    sheet["A2"] = "連絡先電話は090-1111-2222です。"
    workbook.save(path)


def test_highlight_redaction_mode_marks_converted_cell_yellow() -> None:
    # "仮名化＋蛍光ペン" mode: the cell's value is still replaced as normal (the
    # original sensitive text is never kept), but the cell also gets a
    # yellow background fill so a reviewer can see at a glance which cells
    # were actually converted.
    with tempfile.TemporaryDirectory(prefix="excel_highlight_test_") as tmpdir:
        tmp = Path(tmpdir)
        source = tmp / "fixture.xlsx"
        create_simple_xlsx(source)

        processor = ExcelPrivacyProcessor()
        findings = processor.scan(source, options=ProcessingOptions(mode="analysis"))
        for finding in findings:
            finding.enabled = True
        result = processor.convert_with_artifacts(
            source, findings, tmp, options=ProcessingOptions(mode="analysis"), redaction_mode="highlight", write_artifacts=False
        )

        output = load_workbook(result.excel_path)
        cell = output.active["A2"]
        assert_true("090-1111-2222" not in str(cell.value), "Original phone number must not remain")
        assert_true(cell.fill is not None and cell.fill.fgColor.rgb == "FFFFFF00", f"Converted cell should have a yellow fill, got {cell.fill}")


def test_pseudonym_redaction_mode_does_not_apply_highlight() -> None:
    with tempfile.TemporaryDirectory(prefix="excel_highlight_test_") as tmpdir:
        tmp = Path(tmpdir)
        source = tmp / "fixture.xlsx"
        create_simple_xlsx(source)

        processor = ExcelPrivacyProcessor()
        findings = processor.scan(source, options=ProcessingOptions(mode="analysis"))
        for finding in findings:
            finding.enabled = True
        result = processor.convert_with_artifacts(
            source, findings, tmp, options=ProcessingOptions(mode="analysis"), redaction_mode="pseudonym", write_artifacts=False
        )

        output = load_workbook(result.excel_path)
        cell = output.active["A2"]
        assert_true("090-1111-2222" not in str(cell.value), "Original phone number must not remain")
        fill_type = cell.fill.fill_type if cell.fill else None
        assert_true(fill_type is None, f"pseudonym mode must not apply any fill, got fill_type={fill_type}")


if __name__ == "__main__":
    import pytest

    raise SystemExit(pytest.main([__file__, "-v"]))
